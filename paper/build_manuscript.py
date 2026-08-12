"""Build the OpenDC-LCA manuscript DOCX from the reviewed Markdown draft."""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "MANUSCRIPT.md"
OUTPUT = ROOT / "OpenDC-LCA_manuscript.docx"
FIGURE_PNG = ROOT / "figures" / "rendered"

BLUE = "1F4E79"
DARK = "17324D"
MUTED = "60758A"
LIGHT = "EAF1F7"
WHITE = "FFFFFF"

EQUATION_DISPLAY = {
    "1": "E_IT = P_IT u (8760)                                                     (1)",
    "2": "E_facility = E_IT PUE                                                   (2)",
    "3": "I_k,op = EF_k PUE                                                       (3)",
    "4": "I_k,eq = Σ_i q_i ceil(L_s/L_i) (I_k,i^prod + I_k,i^EOL) / L_s          (4)",
    "5": "m_prod,annual = m_0 / L_f + m_loss                                      (5)",
    "6": "I_k,fluid = m_prod,annual I_k^prod + (m_0/L_f) I_k^EOL    and    GHG_direct = m_loss GWP_direct   (6)",
    "7": "z_s = (g_s - g_R) / (g_G - g_R)                                        (7)",
    "8": "I_j,s = I_j,R + (I_j,G - I_j,R) z_s                                    (8)",
    "9": "g* = g_R + (g_G-g_R)(I_b,R-I_a,R)/[(I_a,G-I_a,R)-(I_b,G-I_b,R)]        (9)",
    "10": "g_y^AR5 = [2000 M_CO2 + 28 M_CH4 + 265 M_N2O] 0.45359237 / E_y       (10)",
    "11": "delta_service = min_(j != 2P)(I_j / I_2P) - 1                         (11)",
    "12": "R = 100 (1 - I_alternative / I_baseline)                              (12)",
    "13": "P_c = S_c [(D_c - 1) / 4]                                             (13)",
}


def set_font(run, name="Calibri", size=11, *, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_table_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_fixed_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, before, after in (
        ("Heading 1", 15, 16, 7),
        ("Heading 2", 12.5, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if style_name != "Heading 3" else DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(DARK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("OpenDC-LCA | Manuscript draft")
    set_font(header_run, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_inline_runs(paragraph, text, *, size=10.5):
    text = re.sub(
        r"\\\((.*?)\\\)",
        lambda match: latex_inline_to_plain(match.group(1)),
        text,
    )
    tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, italic=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=size - 0.5, color=DARK)
        else:
            run = paragraph.add_run(token)
            set_font(run, size=size)


def latex_inline_to_plain(text):
    text = text.replace(r"\mathrm{", "").replace("}", "")
    text = text.replace(r"\left", "").replace(r"\right", "")
    text = text.replace(r"\,", " ").replace(r"\times", "×")
    text = text.replace(r"\alpha", "α").replace(r"\beta", "β")
    text = text.replace(r"\sum", "Σ").replace(r"\lceil", "ceil(")
    text = text.replace(r"\rceil", ")")
    text = re.sub(r"_\{([^}]+)\}", r"_\1", text)
    text = re.sub(r"\^\{([^}]+)\}", r"^\1", text)
    return text


def add_title_block(doc, title, author_lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(title)
    set_font(run, size=21, bold=True, color=DARK)
    for line, italic in author_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3 if not italic else 12)
        run = p.add_run(line)
        set_font(run, size=11 if not italic else 9.5, bold=not italic, italic=italic, color=MUTED if italic else DARK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(5)
    rule.paragraph_format.space_after = Pt(8)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_figure(doc, alt_text, relative_svg):
    png_name = Path(relative_svg).stem + ".png"
    png_path = FIGURE_PNG / png_name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(png_path), width=Inches(6.45))
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", f"OpenDC-LCA {alt_text.split('.')[0]}")
    caption = doc.add_paragraph(style="Caption")
    add_inline_runs(caption, alt_text, size=9)


def convert_figures():
    """Render exactly the SVG figures referenced by the main manuscript."""
    FIGURE_PNG.mkdir(parents=True, exist_ok=True)
    used = set(
        re.findall(
            r"\]\(figures/([^)]+\.svg)\)",
            SOURCE.read_text(encoding="utf-8"),
        )
    )
    for filename in used:
        source = ROOT / "figures" / filename
        target = FIGURE_PNG / f"{source.stem}.png"
        if shutil.which("rsvg-convert"):
            subprocess.run(
                [
                    "rsvg-convert",
                    "-w",
                    "2400",
                    "-o",
                    str(target),
                    str(source),
                ],
                check=True,
            )
        else:
            render_svg_with_browser(source, target, width=2400)


def render_svg_with_browser(source: Path, target: Path, *, width: int) -> None:
    """Rasterize an SVG with an installed Chromium browser."""
    browser = next(
        (
            candidate
            for candidate in (
                Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
                Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
                Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
            )
            if candidate.is_file()
        ),
        None,
    )
    if browser is None:
        raise RuntimeError("Figure conversion requires rsvg-convert or Chromium")
    root = ET.parse(source).getroot()
    view_box = root.attrib.get("viewBox", "").split()
    if len(view_box) == 4:
        source_width, source_height = float(view_box[2]), float(view_box[3])
    else:
        source_width = float(root.attrib["width"])
        source_height = float(root.attrib["height"])
    scale = width / source_width
    css_width = round(source_width)
    css_height = round(source_height)
    with tempfile.TemporaryDirectory(prefix="opendc-svg-") as temporary:
        subprocess.run(
            [
                str(browser),
                "--headless=new",
                "--disable-gpu",
                "--hide-scrollbars",
                "--no-first-run",
                f"--force-device-scale-factor={scale}",
                f"--user-data-dir={Path(temporary) / 'profile'}",
                f"--window-size={css_width},{css_height}",
                f"--screenshot={target}",
                source.resolve().as_uri(),
            ],
            check=True,
            timeout=60,
        )
        for _ in range(300):
            if target.is_file() and target.stat().st_size:
                break
            time.sleep(0.1)
        else:
            raise RuntimeError(f"Browser did not render {source.name}")


def add_markdown_table(doc, rows):
    parsed = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    parsed = [parsed[0]] + parsed[2:]
    cols = len(parsed[0])
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = "Table Grid"
    if cols == 5:
        widths = [1300, 1150, 2300, 3000, 1610]
    elif cols == 6:
        widths = [1450, 1750, 1650, 1650, 1850, 1010]
    else:
        widths = [9360 // cols] * cols
        widths[-1] += 9360 - sum(widths)
    for i, row_values in enumerate(parsed):
        for j, value in enumerate(row_values):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline_runs(p, value, size=7.5 if cols >= 6 else 9)
            if i == 0:
                set_cell_shading(cell, BLUE)
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                    run.bold = True
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_table_row_split(row)
    set_fixed_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_markdown(doc, text):
    lines = text.splitlines()
    title = lines[0][2:].strip()
    authors = lines[2].strip("*")
    abstract_index = next(
        i for i, line in enumerate(lines) if line.strip() == "## Abstract"
    )
    title_metadata = [line.strip() for line in lines[4:abstract_index] if line.strip()]
    affiliation_lines = [line for line in title_metadata if not line.startswith("*")]
    draft_note = " ".join(
        line.strip("*") for line in title_metadata if line.startswith("*")
    ) or "Pre-submission draft for coauthor review."
    add_title_block(
        doc,
        title,
        [(authors, False)]
        + [(affiliation, False) for affiliation in affiliation_lines]
        + [(draft_note, True)],
    )

    index = abstract_index
    paragraph_buffer = []

    def flush_paragraph():
        if not paragraph_buffer:
            return
        value = " ".join(line.strip() for line in paragraph_buffer)
        p = doc.add_paragraph()
        add_inline_runs(p, value)
        paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_paragraph(stripped[3:], style="Heading 1")
        elif stripped.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(stripped[4:], style="Heading 2")
        elif stripped.startswith("!["):
            flush_paragraph()
            match = re.match(r"!\[(.*)\]\((.*)\)", stripped)
            if match:
                add_figure(doc, match.group(1), match.group(2))
        elif stripped == r"\[":
            flush_paragraph()
            eq_lines = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                eq_lines.append(lines[index].strip())
                index += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.keep_together = True
            equation_source = " ".join(eq_lines)
            tag_match = re.search(r"\\tag\{(\d+)\}", equation_source)
            equation = EQUATION_DISPLAY.get(
                tag_match.group(1) if tag_match else "",
                latex_inline_to_plain(equation_source),
            )
            run = p.add_run(equation)
            set_font(run, name="Cambria Math", size=10.5, italic=True, color=DARK)
        elif stripped.startswith("|"):
            flush_paragraph()
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(lines[index])
                index += 1
            add_markdown_table(doc, table_rows)
            index -= 1
        elif numbered := re.match(r"^(\d+)\.\s+(.*)", stripped):
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.26)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, f"{numbered.group(1)}. {numbered.group(2)}")
        elif stripped.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
        elif stripped.startswith("**Table "):
            flush_paragraph()
            p = doc.add_paragraph(style="Caption")
            add_inline_runs(p, stripped.strip("*"), size=9)
            p.paragraph_format.keep_with_next = True
        else:
            paragraph_buffer.append(line)
        index += 1
    flush_paragraph()


def main():
    convert_figures()
    doc = Document()
    configure_document(doc)
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    props = doc.core_properties
    props.title = "OpenDC-LCA manuscript"
    props.subject = "Open, evidence-gated LCA for data-center cooling"
    props.author = "Braden Stevens; Pengjiang Xiang; Yimin Chen; Darin Nutter; Han Hu"
    props.keywords = "data center, cooling, LCA, reproducibility, OpenDC-LCA"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
